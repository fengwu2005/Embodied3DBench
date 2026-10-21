import pandas as pd
import numpy as np
import os

# 现在不再重新计算 score，只读取 fengwu_eval_summary*.csv 里每个单元格的 "mean-count" 字符串，
# 直接取 mean 部分作为数值（平均分），再做可视化与六大能力汇总。

# 可根据实际文件名调整
csv_path = './fengwu_eval_summary.csv'
output_image = './score_lines/'
os.makedirs(output_image, exist_ok=True)

# 读取并转置，转置后的第一行作为列名
df = pd.read_csv(csv_path, header=None)
df = df.T
df.loc[0, 0] = 'qa'
df.columns = df.iloc[0]
df = df[1:].reset_index(drop=True)

print(df)

# -------------------------
# 直接提取平均值 (形如 "73.42-12" -> 73.42)
# -------------------------
for qa in df['qa']:
    row_idx = df['qa'] == qa
    for model in df.columns[1:]:
        val = df.loc[row_idx, model].values[0]
        if isinstance(val, str) and '-' in val:
            parts = val.split('-')
            if len(parts) >= 2:
                mean_str = '-'.join(parts[:-1])  # 去掉最后的 count
                try:
                    mean_val = float(mean_str)
                except ValueError:
                    mean_val = np.nan
                df.loc[row_idx, model] = mean_val
            else:
                df.loc[row_idx, model] = np.nan
        else:
            # 如果不是字符串或没有 '-'，尝试直接转为浮点
            try:
                df.loc[row_idx, model] = float(val)
            except (ValueError, TypeError):
                df.loc[row_idx, model] = np.nan

# 保存并对每行取前三名着色
output_xlsx = './fengwu_eval_summary_means.xlsx'
df.to_excel(output_xlsx, index=False)

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
wb = load_workbook(output_xlsx)
ws = wb.active
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=ws.max_column):
    scores = []
    for cell in row:
        try:
            score = float(cell.value)
            scores.append((score, cell))
        except (ValueError, TypeError):
            continue
    scores.sort(key=lambda x: x[0], reverse=True)
    colors = ['FF0000', 'FFA500', 'FFFF00']  # 红 橙 黄
    for i in range(min(3, len(scores))):
        _, cell = scores[i]
        fill = PatternFill(start_color=colors[i], end_color=colors[i], fill_type='solid')
        cell.fill = fill
wb.save(output_xlsx)
print("Mean values with coloring saved.")

six_ability = {
    "Embodied Perception":{
        "Grounding":{
            "2d": [
                "2d_oriented_bbox_qa_detect_category",
                "2d_oriented_bbox_qa_single_part_with_desc",
                "2d_oriented_bbox_qa_single_with_desc",
                "object_point_qa_detect_category_2d",
                "object_point_qa_single_with_desc_2d",
                "single_part_point_2d",
            ],
            "3d": [
                "3d_bbox_qa_single_part_with_desc",
                "obj_bbox_qa_detect_category_3d",
                "obj_bbox_qa_single_with_desc_3d",
                "object_point_qa_detect_category_3d",
                "object_point_qa_single_with_desc_3d",
                "single_part_point_3d",
            ]
        },
        "Correspondence":{
            "2d": [
                "correspondence_pairs_customized_2d",
                "correspondence_pairs_single_object_2d",
            ],
            "3d": [
                "correspondence_pairs_customized_3d",
                "correspondence_pairs_single_object_3d",
            ]
        },
        "Spatial Relation":{
            "all": [
                "spatial_rel_qa_obj2cam_distance",
                "spatial_rel_qa_obj2obj_choice",
                "spatial_rel_qa_obj2obj_distance",
            ]
        }
    },
    "Embodied Manipulation":{
        "Affordance":{
            "2d": [
                "single_func_point_2d",
                "single_func_vector_2d",
            ],
            "3d": [
                "single_func_point_3d",
                "single_func_vector_3d",
            ]
        },
        "Manipulation Trajectory":{
            "2d": ["traj_qa_prediction_2d"],
            "3d": ["traj_qa_prediction_3d"]
        },
        "Grasp Point":{
            "2d": ["grasp_point_2d"],
            "3d": ["grasp_point_3d"]
        }
    }
}

# 六大能力平均值
output_ability_xlsx = './fengwu_eval_summary_six_abilities_means.xlsx'
ability_data = []
for type_a, abilities in six_ability.items():
    for ability, modalities in abilities.items():
        for modality, qa_list in modalities.items():
            row = [type_a, ability, modality]
            for model in df.columns[1:]:
                vals = []
                for qa_task in qa_list:
                    v = df.loc[df['qa'] == qa_task, model].values
                    if len(v) > 0 and not pd.isna(v[0]):
                        vals.append(float(v[0]))
                avg_val = np.mean(vals) if vals else np.nan
                row.append(avg_val)
            ability_data.append(row)

ability_df = pd.DataFrame(ability_data, columns=['type_a', 'ability', 'modality'] + list(df.columns[1:]))
ability_df.to_excel(output_ability_xlsx, index=False)

wb = load_workbook(output_ability_xlsx)
ws = wb.active
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=ws.max_column):
    scores = []
    for cell in row:
        try:
            score = float(cell.value)
            scores.append((score, cell))
        except (ValueError, TypeError):
            continue
    scores.sort(key=lambda x: x[0], reverse=True)
    colors = ['FF0000', 'FFA500', 'FFFF00']
    for i in range(min(3, len(scores))):
        _, cell = scores[i]
        fill = PatternFill(start_color=colors[i], end_color=colors[i], fill_type='solid')
        cell.fill = fill
wb.save(output_ability_xlsx)
